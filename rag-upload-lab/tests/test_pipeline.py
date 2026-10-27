import json
import sys
import types
from pathlib import Path
from contextlib import contextmanager

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app
import state
import worker
from parsing import parse, text_blocks
from chunking import make_chunks


class Tokenizer:
    def encode(self, text, add_special_tokens=True):
        # 多 token 字符用于确保按 token 而非字符长度校验。
        return [1] * (2 * len(text) + 2)


@pytest.fixture
def api(tmp_path, monkeypatch):
    for module in (app, state, worker):
        monkeypatch.setattr(module, "DATA", tmp_path)
    with TestClient(app.app) as c:
        yield c


def upload(api, name="policy.md", content=b"A policy."):
    response = api.post("/api/documents", files={"file": (name, content)})
    assert response.status_code == 202, response.text
    return response.json()


def test_upload_deduplicate_retry_and_recovery(api):
    job = upload(api, "../policy.md")
    assert job["filename"] == "policy.md"
    assert "path" not in job
    assert upload(api)["id"] == job["id"]
    assert len(api.get("/api/documents").json()) == 1
    assert api.post("/api/ask", json={"question": "policy?"}).status_code == 409
    claimed = state.claim()
    assert claimed["id"] == job["id"]
    assert state.claim() is None
    state.recover()
    assert state.get(job["id"])["status"] == "queued"
    state.update(job["id"], status="failed", error="test failure")
    assert api.post(f"/api/documents/{job['id']}/retry").status_code == 202
    assert api.post(f"/api/documents/{job['id']}/retry").status_code == 409


def test_upload_rejects_invalid_inputs(api, monkeypatch):
    assert api.post("/api/documents", files={"file": ("a.exe", b"bad")}).status_code == 415
    assert api.post("/api/documents", files={"file": ("a.txt", b"")}).status_code == 400
    monkeypatch.setattr(app, "MAX_BYTES", 3)
    assert api.post("/api/documents", files={"file": ("a.txt", b"1234")}).status_code == 413
    assert not list((app.DATA / "uploads").iterdir())


def test_chunk_preserves_text_titles_and_budget():
    body = "员工申请软件许可证，需要直属经理审批。" * 80
    blocks = text_blocks("# IT\n\n## 软件\n\n" + body)
    chunks = make_chunks(blocks, Tokenizer(), "doc", "policy.md", "hash")
    assert len(chunks) > 1
    assert all(len(Tokenizer().encode(c["content"])) <= 384 for c in chunks)
    assert all(c["content"].startswith("标题：IT > 软件\n") for c in chunks)
    assert "".join(c["content"].split("\n", 1)[1] for c in chunks) == body
    assert len({c["chunk_id"] for c in chunks}) == len(chunks)


def test_docx_paragraph_and_table_order(tmp_path):
    from docx import Document
    path = tmp_path / "policy.docx"
    doc = Document()
    doc.add_heading("差旅", level=1)
    doc.add_paragraph("以下为住宿标准。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "城市"
    table.cell(0, 1).text = "上限"
    table.cell(1, 0).text = "上海"
    table.cell(1, 1).text = "600元"
    doc.add_paragraph("需要审批。")
    doc.save(path)
    blocks = parse(path)
    assert len(blocks) == 3
    assert blocks[1]["text"] == "城市：上海；上限：600元"
    assert blocks[2]["text"] == "需要审批。"
    assert all(block["heading"] == "差旅" for block in blocks)


def test_docx_title_is_context_not_searchable_chunk(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_heading("IT 政策", level=0)
    doc.add_heading("软件", level=1)
    doc.add_heading("许可证", level=3)
    doc.add_paragraph("需要直属经理审批。")
    doc.add_heading("账号", level=3)
    doc.add_paragraph("通过平台申请。")
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "入口"
    table.cell(1, 0).text = "自助服务平台"
    path = tmp_path / "titled.docx"
    doc.save(path)
    blocks = parse(path)
    assert len(blocks) == 3
    assert blocks[0]["heading"] == "IT 政策 > 软件 > 许可证"
    assert blocks[1]["heading"] == "IT 政策 > 软件 > 账号"
    assert blocks[2]["heading"] == blocks[1]["heading"]
    assert all(block["text"] != "IT 政策" for block in blocks)


def test_word_reindex_queues_fresh_artifacts_and_preserves_old(api):
    from reindex_document import enqueue
    job = upload(api, "policy.docx")
    state.update(job["id"], status="ready", store_id="vs_old", total=4, done=4)
    new = enqueue(job["id"])
    assert new["id"] != job["id"]
    assert state.get(job["id"])["status"] == "ready"
    pending = state.get(new["id"])
    assert pending["status"] == "queued"
    assert pending["store_id"] is None
    assert pending["total"] == 0
    assert enqueue(job["id"])["id"] == new["id"]


def test_failed_batch_retry_reuses_store_and_ids(api, monkeypatch):
    job = upload(api, content="\n\n".join(f"政策条目 {i}" for i in range(18)).encode())
    fake_transformers = types.SimpleNamespace(AutoTokenizer=types.SimpleNamespace(
        from_pretrained=lambda *a, **k: Tokenizer()))
    monkeypatch.setitem(sys.modules, "transformers", fake_transformers)

    @contextmanager
    def fake_client():
        yield None
    monkeypatch.setattr(worker, "client", fake_client)
    records = {"creates": 0, "ids": [], "fail": True}

    def fake_call(c, method, path, body=None):
        if path == "/vector_stores":
            records["creates"] += 1
            return {"id": "vs_test"}
        if path == "/vector_stores/vs_test":
            return {"id": "vs_test"}
        if path == "/embeddings":
            return {"data": [{"index": i, "embedding": [0.1] * 512} for i, _ in enumerate(body["input"])]}
        if path == "/vector-io/insert":
            assert state.get(job["id"])["status"] != "ready"
            if len(body["chunks"]) == 2 and records["fail"]:
                raise RuntimeError("simulated network failure")
            records["ids"].extend(chunk["chunk_id"] for chunk in body["chunks"])
            return None
        raise AssertionError(path)
    monkeypatch.setattr(worker, "call", fake_call)
    state.claim()
    with pytest.raises(RuntimeError, match="simulated"):
        worker.process(job["id"])
    assert state.get(job["id"])["done"] == 16
    assert api.post("/api/ask", json={"question": "test"}).status_code == 409
    state.update(job["id"], status="failed")
    api.post(f"/api/documents/{job['id']}/retry")
    state.claim()
    records["fail"] = False
    worker.process(job["id"])
    assert records["creates"] == 1
    assert records["ids"][:16] == records["ids"][16:32]
    assert state.get(job["id"])["status"] == "ready"
    chunks = api.get(f"/api/documents/{job['id']}/chunks").json()
    assert len(chunks) == 18


def test_ready_only_and_latest_filename_selection(api, monkeypatch):
    first = upload(api, content=b"old")
    second = upload(api, content=b"new")
    state.update(first["id"], status="ready", store_id="vs_old")
    state.update(second["id"], status="ready", store_id="vs_new")
    calls = []

    @contextmanager
    def fake_client():
        yield None
    monkeypatch.setattr(app, "client", fake_client)

    def fake_call(c, method, path, body=None):
        calls.append(body["vector_store_id"])
        return {"chunks": [], "scores": []}
    monkeypatch.setattr(app, "call", fake_call)
    assert api.post("/api/ask", json={"question": "test"}).status_code == 200
    assert calls == ["vs_new"]
    assert api.post("/api/ask", json={"question": "test", "document_ids": ["missing"]}).status_code == 409
    assert api.post("/api/ask", json={"question": "test", "document_ids": []}).status_code == 400


def test_search_returns_evidence_without_generation(api, monkeypatch):
    job = upload(api)
    state.update(job["id"], status="ready", store_id="vs_one")

    @contextmanager
    def fake_client():
        yield None
    monkeypatch.setattr(app, "client", fake_client)
    def fake_call(c, method, path, body=None):
        assert path == "/vector-io/query"
        return {"chunks": [{"chunk_id": "c1", "content": "Policy text",
                            "metadata": {"document_id": job["id"], "source": "policy.md", "block": 1}}],
                "scores": [2.0]}
    monkeypatch.setattr(app, "call", fake_call)
    response = api.post("/api/search", json={"question": "policy"})
    assert response.status_code == 200
    assert response.json()["sources"][0]["document_id"] == job["id"]
    assert "answer" not in response.json()
